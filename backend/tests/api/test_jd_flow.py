from io import BytesIO
from pathlib import Path
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import SecretStr
from sqlalchemy import select

from kerui_recruit.core.settings import Settings
from kerui_recruit.db.models import JdRevision, MatchRun
from kerui_recruit.runtime import create_runtime_app


@pytest.fixture
def client(tmp_path: Path):
    settings = Settings(
        data_root=tmp_path / "data",
        session_token=SecretStr("launch-token"),
    )
    app = create_runtime_app(settings)
    with TestClient(app) as test_client:
        yield test_client


def _headers() -> dict[str, str]:
    return {"X-Kerui-Session": "launch-token"}


def _wait_jd_ready(client: TestClient, revision_id: str) -> None:
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        response = client.get("/api/jd", headers=_headers())
        assert response.status_code == 200
        revision = next((item for item in response.json() if item["revision_id"] == revision_id), None)
        if revision and revision["status"] == "READY":
            assert revision["jd_status"] == "OPEN"
            pending = client.app.state.services.index_sync_service.status()["items"]
            if not any(item["entity_type"] == "jd" and item["entity_id"] == revision["jd_id"] for item in pending):
                return
        time.sleep(.05)
    raise AssertionError("Local JD pipeline did not produce an OPEN, READY revision")


def test_jd_import_and_match_round_trip(client: TestClient) -> None:
    imported = client.post(
        "/api/jd/import",
        json={"company": "某金融", "title": "Java 后端", "source_text": "Java 3年 本科 金融"},
        headers=_headers(),
    )
    assert imported.status_code == 202
    body = imported.json()
    assert body["jd_id"]
    _wait_jd_ready(client, body["revision_id"])

    match = client.post(
        "/api/match/jd",
        json={"revision_id": body["revision_id"], "limit": 20},
        headers=_headers(),
    )
    assert match.status_code == 200


def test_batch_match_multiple_jds(client: TestClient) -> None:
    first = client.post(
        "/api/jd/import",
        json={"company": "某金融", "title": "Java 后端", "source_text": "Java 3年 本科 金融"},
        headers=_headers(),
    ).json()
    second = client.post(
        "/api/jd/import",
        json={"company": "某科技", "title": "算法工程师", "source_text": "Python 算法 大模型"},
        headers=_headers(),
    ).json()
    _wait_jd_ready(client, first["revision_id"])
    _wait_jd_ready(client, second["revision_id"])

    batch = client.post(
        "/api/match/batch",
        json={"revision_ids": [first["revision_id"], second["revision_id"]], "limit": 20},
        headers=_headers(),
    )

    assert batch.status_code == 200
    results = batch.json()["results"]
    assert len(results) == 2
    assert {r["revision_id"] for r in results} == {first["revision_id"], second["revision_id"]}
    assert all(r["run_id"] for r in results)


def test_match_persists_requested_mode(client: TestClient) -> None:
    imported = client.post(
        "/api/jd/import",
        json={"company": "某金融", "title": "Java 后端", "source_text": "Java 3年 本科 金融"},
        headers=_headers(),
    ).json()
    _wait_jd_ready(client, imported["revision_id"])

    batch = client.post(
        "/api/match/batch",
        json={"revision_ids": [imported["revision_id"]], "limit": 20, "mode": "vector"},
        headers=_headers(),
    )
    assert batch.status_code == 200
    run_id = batch.json()["results"][0]["run_id"]

    factory = client.app.state.services.session_factory
    with factory() as session:
        run = session.get(MatchRun, run_id)
        assert run is not None
        assert run.mode == "vector"


def test_match_jd_rejects_invalid_mode(client: TestClient) -> None:
    imported = client.post(
        "/api/jd/import",
        json={"company": "某金融", "title": "Java 后端", "source_text": "Java 3年 本科 金融"},
        headers=_headers(),
    ).json()
    _wait_jd_ready(client, imported["revision_id"])

    match = client.post(
        "/api/match/jd",
        json={"revision_id": imported["revision_id"], "limit": 20, "mode": "semantic"},
        headers=_headers(),
    )
    assert match.status_code == 422


def _docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_import_jd_from_word_and_excel(client: TestClient) -> None:
    docx = _docx_bytes("Java 后端工程师 3年 Java 本科 金融")
    imported_docx = client.post(
        "/api/jd/import-file",
        files={"file": ("jd.docx", docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"company": "某金融", "title": "Java 后端"},
        headers=_headers(),
    )
    assert imported_docx.status_code == 202
    assert imported_docx.json()["jd_id"]

    xlsx = _xlsx_bytes([["岗位", "算法工程师"], ["要求", "Python 算法 大模型"]])
    imported_xlsx = client.post(
        "/api/jd/import-file",
        files={"file": ("jd.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"company": "某科技", "title": "算法工程师"},
        headers=_headers(),
    )
    assert imported_xlsx.status_code == 202
    assert imported_xlsx.json()["jd_id"]


def test_update_jd_candidate_profile_persists_and_records_override(client: TestClient) -> None:
    imported = client.post(
        "/api/jd/import",
        json={"company": "某金融", "title": "Java 后端", "source_text": "Java 3年 本科 金融"},
        headers=_headers(),
    ).json()
    _wait_jd_ready(client, imported["revision_id"])

    edit = client.put(
        f"/api/jd/{imported['jd_id']}/field",
        json={"field": "candidate_profile", "value": "手工画像"},
        headers=_headers(),
    )
    assert edit.status_code == 200
    assert edit.json()["value"] == "手工画像"

    factory = client.app.state.services.session_factory
    with factory() as session:
        revision = session.get(JdRevision, imported["revision_id"])
        assert revision.parsed_data["candidate_profile"] == "手工画像"
        assert revision.manual_overrides.get("candidate_profile") == "手工画像"


def test_jd_profile_stream_reports_stages_then_result(client: TestClient) -> None:
    """`Accept: text/event-stream` 时先推阶段再推结果（问题 #6 的另一半：等待期要有反馈）。

    只有转圈看不出「快好了」和「刚开始第二次模型调用」，而后者意味着还要再等一个完整轮次。
    """
    imported = client.post(
        "/api/jd/import",
        json={"company": "某金融", "title": "Java 后端", "source_text": "Java 3年 本科 金融"},
        headers=_headers(),
    )
    assert imported.status_code in (200, 202), imported.text
    jd_id = imported.json()["jd_id"]

    class _StreamingBackfill:
        async def regenerate_jd_profile(self, jd_id: str, instruction=None, *, on_stage=None):
            assert on_stage is not None, "流式路径必须把阶段回调传下去"
            on_stage("loading")
            on_stage("draft")
            on_stage("repair")
            return {"generated": True, "summary": "三年 Java 后端经验。",
                    "points": [{"text": "三年 Java 后端经验。", "evidence_paths": []}],
                    "compact": "三年 Java 后端经验。"}

    services = client.app.state.services
    original = services.backfill_service
    object.__setattr__(services, "backfill_service", _StreamingBackfill())
    try:
        response = client.post(
            f"/api/jd/{jd_id}/regen-profile",
            json={"instruction": ""},
            headers={**_headers(), "Accept": "text/event-stream"},
        )
    finally:
        object.__setattr__(services, "backfill_service", original)

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/event-stream")
    body = response.text
    for stage in ("loading", "draft", "repair"):
        assert f'"stage": "{stage}"' in body, body
    # 阶段必须带可展示的中文文案（前端不自己拼文案，见 api/profile_stream.py）。
    assert '"message": "正在生成画像初稿…"' in body, body
    assert "event: result" in body
    assert '"generated": true' in body


def test_jd_profile_stream_returns_errors_as_events(client: TestClient) -> None:
    """流已开始就改不了状态码，所以失败必须走 `error` 事件，否则客户端只看到一条断流。"""
    from kerui_recruit.providers.errors import ProviderError

    imported = client.post(
        "/api/jd/import",
        json={"company": "某金融", "title": "Java 后端", "source_text": "Java 3年 本科 金融"},
        headers=_headers(),
    )
    assert imported.status_code in (200, 202), imported.text
    jd_id = imported.json()["jd_id"]

    class _UnavailableBackfill:
        async def regenerate_jd_profile(self, jd_id: str, instruction=None, *, on_stage=None):
            raise ProviderError(code="E_AI_NO_PROVIDER", retryable=True,
                                user_message="没有可用的 AI 服务")

    services = client.app.state.services
    original = services.backfill_service
    object.__setattr__(services, "backfill_service", _UnavailableBackfill())
    try:
        response = client.post(
            f"/api/jd/{jd_id}/regen-profile",
            json={"instruction": ""},
            headers={**_headers(), "Accept": "text/event-stream"},
        )
    finally:
        object.__setattr__(services, "backfill_service", original)

    # HTTP 200 是 SSE 的固有形态（响应头先发出）；错误码在事件体里，与同步路径逐字一致。
    assert response.status_code == 200, response.text
    assert "event: error" in response.text
    assert '"code": "E_AI_NO_PROVIDER"' in response.text
    assert '"status": 503' in response.text
    assert '"message": "没有可用的 AI 服务"' in response.text


def test_jd_profile_regeneration_surfaces_provider_unavailable(client: TestClient) -> None:
    """岗位画像重生成遇到供应商不可用要回 502/503，不能压成 500 内部错误（实测同候选人侧）。"""
    from kerui_recruit.providers.errors import ProviderError

    imported = client.post(
        "/api/jd/import",
        json={"company": "某金融", "title": "Java 后端", "source_text": "Java 3年 本科 金融"},
        headers=_headers(),
    )
    assert imported.status_code in (200, 202), imported.text
    jd_id = imported.json()["jd_id"]

    class _UnavailableBackfill:
        async def regenerate_jd_profile(self, jd_id: str, instruction=None, *, on_stage=None):
            raise ProviderError(code="E_AI_NO_PROVIDER", retryable=True,
                                user_message="没有可用的 AI 服务")

    services = client.app.state.services
    original = services.backfill_service
    object.__setattr__(services, "backfill_service", _UnavailableBackfill())
    try:
        response = client.post(f"/api/jd/{jd_id}/regen-profile",
                               json={"instruction": "一句话概括"}, headers=_headers())
    finally:
        object.__setattr__(services, "backfill_service", original)
    assert response.status_code == 503, response.text
    assert response.json()["code"] == "E_AI_NO_PROVIDER"


def test_profile_edit_drives_min_years_and_keeps_it_when_unmentioned(client: TestClient) -> None:
    """画像里的年限要落到 min_years；画像没提年限时不能被编辑抹掉。

    年限不在 exact_constraints 的 kind 里（它是独立的硬窗口），所以必须单独接进画像链路，
    否则「在画像里写 5 年以上」不会有任何筛选效果。
    """
    imported = client.post(
        "/api/jd/import",
        json={"company": "某金融", "title": "Java 后端", "source_text": "Java 3年 本科 金融"},
        headers=_headers(),
    ).json()
    _wait_jd_ready(client, imported["revision_id"])
    factory = client.app.state.services.session_factory

    edit = client.put(
        f"/api/jd/{imported['jd_id']}/parsed",
        json={"parsed_data": {"candidate_profile": "要求 5 年以上后端经验，熟悉 Java"}},
        headers=_headers(),
    )
    assert edit.status_code == 200
    with factory() as session:
        revision = session.get(JdRevision, imported["revision_id"])
        assert revision.min_years == 5
        assert revision.parsed_data["min_years"] == 5

    # 后续画像不再提年限 → 保留上一版的 5，而不是抹成 NULL。
    keep = client.put(
        f"/api/jd/{imported['jd_id']}/parsed",
        json={"parsed_data": {"candidate_profile": "熟悉 Java 与微服务"}},
        headers=_headers(),
    )
    assert keep.status_code == 200
    with factory() as session:
        assert session.get(JdRevision, imported["revision_id"]).min_years == 5

    # 明确「经验不限」→ 清掉年限窗口。
    clear = client.put(
        f"/api/jd/{imported['jd_id']}/parsed",
        json={"parsed_data": {"candidate_profile": "熟悉 Java，经验不限"}},
        headers=_headers(),
    )
    assert clear.status_code == 200
    with factory() as session:
        assert session.get(JdRevision, imported["revision_id"]).min_years is None


@pytest.mark.asyncio
async def test_parse_constraints_prefers_model_years_with_rule_fallback():
    """年限：模型优先；模型漏判年限或整个不可用时用正则兜底（画像写明的年限不能静默失效）。"""
    from types import SimpleNamespace

    from kerui_recruit.api.jd import ParseConstraintsRequest, parse_constraints
    from kerui_recruit.jd.profile_constraints import ProfileRequirements

    class _Generator:
        def __init__(self, result):
            self._result = result

        async def parse_constraints(self, source_text, **kwargs):
            return self._result

    def _request(generator):
        services = SimpleNamespace(backfill_service=SimpleNamespace(jd_generator=generator))
        return SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(services=services)))

    # 模型给出年限 → 以模型为准（即便正则能从同一段文本里抽到别的数字）。
    model = ProfileRequirements(constraints=[], min_years=8.0, years_stated=True)
    response = await parse_constraints(
        ParseConstraintsRequest(source_text="5 年以上经验"), _request(_Generator(model))
    )
    assert (response.min_years, response.years_stated) == (8.0, True)

    # 模型抽到了硬条件但漏判年限 → 年限单独走一次规则兜底。
    silent = ProfileRequirements(
        constraints=[{"kind": "skill", "operator": "OR", "alternatives": ["Java"],
                      "strength": "PLUS", "source": "inferred", "source_text": "熟悉 Java"}],
        min_years=None,
        years_stated=False,
    )
    response = await parse_constraints(
        ParseConstraintsRequest(source_text="熟悉 Java，5 年以上经验"), _request(_Generator(silent))
    )
    assert (response.min_years, response.years_stated) == (5.0, True)

    # 模型不可用 → 整条规则兜底；「经验不限」是「明确提及 + 无年限」，与「没提」不是一回事。
    response = await parse_constraints(
        ParseConstraintsRequest(source_text="熟悉 Java，经验不限"), _request(None)
    )
    assert (response.min_years, response.years_stated) == (None, True)


def test_candidate_soft_delete_is_rejected(client: TestClient) -> None:
    resp = client.post(
        "/api/soft-delete",
        json={"entity_type": "candidate", "entity_id": "any-candidate"},
        headers=_headers(),
    )
    assert resp.status_code == 400

    restore = client.post(
        "/api/soft-delete/restore",
        json={"entity_type": "candidate", "entity_id": "any-candidate"},
        headers=_headers(),
    )
    assert restore.status_code == 400
