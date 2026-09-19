from __future__ import annotations

import importlib
import platform
import re
import subprocess
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml.ns import qn

from kerui_recruit.resumes.quality import (
    DOMINANT_FRAGMENT_RATIO,
    IMAGE_HEAVY_COVERAGE,
    MIN_MEANINGFUL_CHARS,
    analyze_text,
)

_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_RE = re.compile(r"(?<!\d)1[3-9](?:[\s-]?\d){9}(?!\d)")

# 无法确认姓名时的中性占位值：绝不用文件名冒充姓名。
# build_display_name / ingest / 解析持久化共用同一口径。
UNKNOWN_NAME = "未知"

# 姓名线索只在前若干行里找：简历中部和末尾的短行更可能是正文或其他人的信息。
_NAME_HEAD_LINES = 20
# 姓名行一定很短；超长行（声明、段落）只在联系方式同行时按相邻位置找姓名。
_NAME_LINE_MAX_CHARS = 60
# 与手机号/邮箱同行的行若超过这个长度，基本是正文段落而不是个人信息行。
_CONTACT_LINE_MAX_CHARS = 80

_NAME_LABEL_RE = re.compile(
    r"(?:姓\s*名|名字|(?<![a-zA-Z])Name)\s*[:：]?\s*([^\s|｜,，;；、/／]{2,30})",
    re.IGNORECASE,
)
_SEGMENT_SPLIT_RE = re.compile(r"[|｜,，;；、/／\t ]+")
# 全文扫称谓（"张三先生"/"Mr. Smith"）；要求左侧是行首或分隔符，
# 避免把「推荐人是王女士」这种句子切出半个名字。
_HONORIFIC_RE = re.compile(
    r"(?:^|[\s|｜,，;；、(（【])"
    r"((?:Mr|Ms|Mrs|Miss)\.?\s+[A-Z][a-zA-Z'’-]+(?:\s+[A-Z][a-zA-Z'’-]+)?"
    r"|[\u4e00-\u9fa5]{1,4}(?:先生|女士|小姐))"
)
_HONORIFIC_CN_FULL_RE = re.compile(r"([\u4e00-\u9fa5]{1,4})(?:先生|女士|小姐)")
_HONORIFIC_EN_FULL_RE = re.compile(
    r"(?:Mr|Ms|Mrs|Miss)\.?\s+([A-Z][a-zA-Z'’-]+(?:\s+[A-Z][a-zA-Z'’-]+)?)"
)
# 文档标题里的姓名包装：「张三的简历」「个人简历」。必须带「的」或「个人」，
# 否则「后端简历」这类「岗位+简历」会被削成「后端」当成姓名。
_RESUME_WRAPPER_RE = re.compile(
    r"^(?:个人)?(?:简历|履历)[-—_·\s]*|[-—_·\s]*(?:的个人|的|个人)(?:简历|履历)$"
)
_CN_NAME_RE = re.compile(r"[\u4e00-\u9fa5]{1,4}(?:·[\u4e00-\u9fa5]{1,8})*")
_EN_NAME_RE = re.compile(r"[A-Z][a-zA-Z'’-]*(?:\s+[A-Z][a-zA-Z'’-]*)+")
_EMAIL_LOCAL_RE = re.compile(r"^([a-zA-Z]{2,12})[._-]([a-zA-Z]{2,12})$")
# 文件名/长行里的姓名片段：按非姓名可用的字符切开。
_FILENAME_TOKEN_RE = re.compile(r"[^\u4e00-\u9fa5a-zA-Z·]+")

# 姓氏表：《百家姓》常用单姓 + 复姓。中文姓名必须以此开头——
# 「个人优势」「现居台北」「美团金服」这类 2~4 字短行形状与姓名一致，
# 只有姓氏锚定才能挡住（真实语料里逐个踩过）。
_CN_SURNAMES = frozenset(
    "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜戚谢邹喻柏水窦章"
    "云苏潘葛奚范彭郎鲁韦昌马苗凤花方俞任袁柳酆鲍史唐费廉岑薛雷贺倪汤滕殷罗毕郝邬安常"
    "乐于时傅皮卞齐康伍余元卜顾孟平黄和穆萧尹姚邵湛汪祁毛禹狄米贝明臧计伏成戴谈宋茅庞"
    "熊纪舒屈项祝董梁杜阮蓝闵席季麻强贾路娄危江童颜郭梅盛林刁钟徐邱骆高夏蔡田樊胡凌霍"
    "虞万支柯昝管卢莫经房裘缪干解应宗丁宣贲邓郁单杭洪包诸左石崔吉钮龚程嵇邢滑裴陆荣翁"
    "荀羊於惠甄曲家封芮羿储靳汲邴糜松井段富巫乌焦巴弓牧隗山谷车侯宓蓬全郗班仰秋仲伊宫"
    "宁仇栾暴甘钭历戎祖武符刘景詹束龙叶幸司韶郜黎蓟薄印宿白怀蒲邰从鄂索咸籍赖卓蔺屠蒙"
    "池乔阴胥能苍双闻莘党翟谭贡劳逄姬申扶堵冉宰郦雍却璩桑桂濮牛寿通边扈燕冀郏浦尚农"
    "温别庄晏柴瞿阎充慕连茹习宦艾鱼容向古易慎戈廖庾终暨居衡步都耿满弘匡国文寇广禄阙东"
    "欧殳沃利蔚越夔隆师巩厍聂晁勾敖融冷訾辛阚那简饶空曾毋沙乜养鞠须丰巢关蒯相查后荆红"
    "游竺权逯盖益桓公"
    # 补常见/易漏单姓（用真实候选人姓名逐个校验过）
    "付代闫丘麦南肖兰区涂原初尉伦商芦候仝朴"
)
_CN_COMPOUND_SURNAMES = frozenset({
    "欧阳", "太史", "端木", "上官", "司马", "东方", "独孤", "南宫", "万俟", "闻人",
    "夏侯", "诸葛", "尉迟", "公羊", "赫连", "澹台", "皇甫", "宗政", "濮阳", "公冶",
    "太叔", "申屠", "公孙", "慕容", "仲孙", "钟离", "长孙", "宇文", "司徒", "鲜于",
    "司空", "闾丘", "子车", "亓官", "司寇", "巫马", "公西", "颛孙", "壤驷", "公良",
    "漆雕", "乐正", "宰父", "谷梁", "拓跋", "夹谷", "轩辕", "令狐", "段干", "百里",
    "呼延", "东郭", "南门", "羊舌", "微生", "梁丘", "左丘", "西门", "第五", "公乘",
    "贯丘", "南荣", "东里", "仲长", "即墨", "达奚", "褚师",
})

# 学历/性别/城市/栏目名等短词：形状像姓名，但绝不能当姓名。
_NOT_NAME_WORDS = frozenset({
    "姓名", "名字", "个人", "性别", "年龄", "电话", "邮箱", "手机", "学历",
    "专业", "籍贯", "民族", "住址", "地址", "生日", "婚否",
    "本科", "硕士", "博士", "大专", "中专", "高中", "研究生", "学士", "学位",
    "男", "女", "男性", "女性",
    "北京", "上海", "深圳", "广州", "杭州", "成都", "南京", "武汉", "西安",
    "苏州", "天津", "重庆", "长沙", "郑州", "青岛", "厦门", "合肥", "福州",
    "济南", "东莞", "香港", "澳门", "台北", "海外", "应届", "在读", "在职",
    # 真实语料里被误判过的短词：连接词/公文用语/岗位通称。
    "否则", "以上", "以下", "其他", "备注", "说明", "声明", "注意", "如有",
    "请勿", "实习", "兼职", "全职", "远程", "面议",
})
# 栏目名/职位名：出现在候选行里即否决。
_NOT_NAME_MARKERS = (
    "简历", "履历", "个人信息", "基本信息", "教育", "经历", "技能", "评价",
    "背景", "项目", "联系", "应聘", "求职", "招聘", "网站", "简介", "概述",
    "总结", "荣誉", "奖项", "证书", "实习", "校园", "研究", "特长", "兴趣",
    "作品", "专利", "论文", "自荐", "意向", "期望", "附件", "备注",
    "本科", "硕士", "博士", "大专", "中专", "研究生", "学士", "学位", "学历",
    "开发", "前端", "后端", "全栈", "运维", "测试", "算法", "数据", "产品",
    "设计", "项目", "管理", "工程", "架构", "安全",
    "工程师", "经理", "主管", "总监", "助理", "专员", "顾问", "设计师",
    "架构师", "分析师", "运营", "销售", "会计", "律师", "医生", "教师",
    "讲师", "教授", "董事长", "总裁", "组长", "队长", "班长", "主任", "部长",
)
_ORG_SUFFIXES = (
    "大学", "学院", "学校", "中学", "小学", "高中", "公司", "集团", "银行",
    "医院", "研究所", "研究院", "分行", "支行", "事业部", "工作室",
    "科技", "网络", "技术", "信息", "软件", "电子", "有限", "中心",
)
_EN_NOT_NAME_WORDS = frozenset({
    "resume", "curriculum", "vitae", "personal", "information", "profile",
    "education", "experience", "skills", "skill", "contact", "summary",
    "objective", "project", "projects", "name", "phone", "email", "address",
    "developer", "engineer", "manager", "director", "specialist", "consultant",
    "analyst", "architect", "designer", "leader", "officer", "senior", "junior",
    "intern", "company", "university", "college", "school", "introduction",
    "professional", "confidential", "internal", "draft",
    # 技术词/职位词：形状同样是「两个首字母大写的单词」，但不是姓名。
    "python", "java", "javascript", "typescript", "golang", "node", "react",
    "vue", "sql", "data", "cloud", "backend", "frontend", "fullstack", "stack",
    "devops", "mobile", "android", "ios", "web", "software", "system", "systems",
    "network", "security", "marketing", "finance", "sales", "product", "design",
    "operations", "research", "quality", "test", "testing", "engineering",
    "development", "science", "artificial", "intelligence", "machine",
    "learning", "analysis", "business", "management", "technology", "solution",
    "solutions", "platform", "architecture",
})
_EN_EMAIL_STOPWORDS = frozenset({
    "hr", "admin", "info", "contact", "mail", "email", "job", "jobs",
    "recruit", "recruiting", "office", "service", "support", "hello", "team",
})
_PLACEHOLDER_NAMES = frozenset({
    "未知", "未知姓名", "待确认", "待识别", "待补充", "无", "暂无",
    "unknown", "n/a", "na", "none", "null", "-", "--",
})


class LegacyDocConversionError(RuntimeError):
    code = "E_DOC_CONVERTER_UNAVAILABLE"


@dataclass(frozen=True, slots=True)
class ExtractedContact:
    email: str | None
    phone: str | None


def extract_contact(text: str) -> ExtractedContact:
    """Extract a best-effort email and mainland mobile number from raw text."""
    email_match = _EMAIL_RE.search(text)
    phone_match = _PHONE_RE.search(text)
    phone = phone_match.group(0) if phone_match else None
    if phone:
        phone = re.sub(r"\D", "", phone)
    return ExtractedContact(
        email=email_match.group(0) if email_match else None,
        phone=phone,
    )


def is_placeholder_name(name: str | None) -> bool:
    """占位名（未知/空值）不参与「仅姓名一致」的疑似重复判断。"""
    value = " ".join((name or "").split())
    return not value or value.casefold() in _PLACEHOLDER_NAMES


def _has_surname_prefix(value: str) -> bool:
    """中文姓名必须以姓氏开头（复姓看前两个字）。"""
    if value[:2] in _CN_COMPOUND_SURNAMES:
        return True
    return value[0] in _CN_SURNAMES


def _is_chinese_name(value: str, *, single_char_ok: bool = False) -> bool:
    if not _CN_NAME_RE.fullmatch(value):
        return False
    # 单字只接受「李女士」这类带称谓的写法，避免把正文里的单字行当姓名。
    if len(value) < 2 and not single_char_ok:
        return False
    if value in _NOT_NAME_WORDS:
        return False
    if any(marker in value for marker in _NOT_NAME_MARKERS):
        return False
    if value.endswith(_ORG_SUFFIXES):
        return False
    # 少数民族姓名（含间隔号）不做姓氏约束。
    return "·" in value or _has_surname_prefix(value)


def _is_english_name(value: str) -> bool:
    # 要求两段以上："John Smith" 可信；单个单词（"Python"/"Jessica"）歧义太大。
    if not _EN_NAME_RE.fullmatch(value):
        return False
    return not any(word.casefold() in _EN_NOT_NAME_WORDS for word in value.split())


def _name_from_segment(segment: str) -> str | None:
    """判断一个片段是否是姓名；是则返回去掉标签/称谓/标题包装后的姓名。"""
    value = segment.strip().strip("•·-—_\u3000 ")
    value = _RESUME_WRAPPER_RE.sub("", value).strip()
    value = " ".join(value.split())
    if not value or len(value) > 40:
        return None

    single_char_ok = False
    honorific_cn = _HONORIFIC_CN_FULL_RE.fullmatch(value)
    honorific_en = _HONORIFIC_EN_FULL_RE.fullmatch(value)
    if honorific_cn:
        # 「李女士」只知道姓，也算一条姓名线索（比占位名有信息量）。
        value, single_char_ok = honorific_cn.group(1), True
    elif honorific_en:
        value = honorific_en.group(1)
    value = value.strip(" .,，、")

    if _is_chinese_name(value, single_char_ok=single_char_ok):
        return value
    if _is_english_name(value):
        return value
    return None


def _name_from_email(email: str) -> str | None:
    """邮箱前缀里的拼音姓名：「zhang.san@x.com」→「Zhang San」。

    只接受 first.last 两段式纯字母，且排除 hr/admin、Java 这类通用前缀。
    """
    local = email.split("@", 1)[0]
    match = _EMAIL_LOCAL_RE.fullmatch(local)
    if match is None:
        return None
    parts = [part for part in match.groups() if part.casefold() not in _EN_EMAIL_STOPWORDS]
    if len(parts) != 2 or any(part.casefold() in _EN_NOT_NAME_WORDS for part in parts):
        return None
    return " ".join(part.capitalize() for part in parts)


def name_from_filename(filename: str | None) -> str | None:
    """从文件名里取出姓名，只在正文完全没有姓名线索时兜底。

    「张三.docx」「MU Congshan.pdf」直接用整名；「李勇-38岁.docx」
    「IT基础主管-陈军-MGA.pdf」「【供应链产品总监】曹瑞峰 13810161818.pdf」
    这类带岗位/年份/联系方式的，按姓氏锚定挑出姓名片段。
    """
    if not filename:
        return None
    stem = Path(filename.replace("\\", "/")).stem
    whole = _name_from_segment(stem)
    if whole:
        return whole
    for token in _FILENAME_TOKEN_RE.split(stem):
        name = _name_from_segment(token)
        if name:
            return name
    return None


def extract_name(text: str) -> str | None:
    """从简历原文里尽力找出候选人姓名，找不到返回 None。

    模型判空不代表原文没有姓名：姓名栏缺失时还有「张三先生」「Jessica Chen」
    这类称谓/英文名写法。这里按证据强度分层兜底（标签 > 联系方式同行 > 首行 >
    称谓 > 邮箱前缀），只返回最有把握的一个，不做任何拼接或臆造。
    """
    if not text:
        return None
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return None
    head = lines[:_NAME_HEAD_LINES]

    # 1) 显式姓名栏：姓名 / 名字 / Name（"姓名：张三"、"Name: John Smith"）
    for match in _NAME_LABEL_RE.finditer(text):
        name = _name_from_segment(match.group(1))
        if name:
            return name

    # 2) 与手机号/邮箱同一行：这一行是个人信息行，姓名就在联系方式旁边。
    # 只取紧邻联系方式的片段——长段落里同样会出现「否则」这类两字词，
    # 一旦按片段扫全文就会当成姓名。
    for line in head:
        if len(line) > _CONTACT_LINE_MAX_CHARS:
            continue
        if not (_PHONE_RE.search(line) or _EMAIL_RE.search(line)):
            continue
        segments = [segment for segment in _SEGMENT_SPLIT_RE.split(line) if segment]
        for index, segment in enumerate(segments):
            if not (_PHONE_RE.search(segment) or _EMAIL_RE.search(segment)):
                continue
            for window in (
                segments[max(0, index - 2):index],
                segments[index + 1:index + 3],
                segments[max(0, index - 1):index],
                segments[index + 1:index + 2],
            ):
                name = _name_from_segment(" ".join(window))
                if name:
                    return name

    # 3) 顶部几行：绝大多数简历第一行就是姓名（手机号/邮箱排在姓名前面时在第二三行）。
    # 这一层是"盲猜"，只认「整行就是姓名」，不切片段——否则会把「否则」「开发」
    # 这类两字词当姓名。
    for line in head[:3]:
        if len(line) > _NAME_LINE_MAX_CHARS:
            continue
        name = _name_from_segment(line)
        if name:
            return name

    # 4) 称谓：张三先生 / 李女士 / Mr. Smith
    for line in head:
        match = _HONORIFIC_RE.search(line)
        if match:
            name = _name_from_segment(match.group(1))
            if name:
                return name

    # 5) 邮箱前缀里的拼音姓名
    for line in lines:
        match = _EMAIL_RE.search(line)
        if match:
            name = _name_from_email(match.group(0))
            if name:
                return name
    return None


@dataclass(frozen=True, slots=True)
class PageAssessment:
    page_index: int
    text: str
    needs_ocr: bool
    reason: str
    valid_char_count: int
    repeated_ratio: float
    dominant_ratio: float
    image_coverage: float


@dataclass(frozen=True, slots=True)
class ExtractedText:
    text: str
    page_count: int
    requires_ocr: bool
    pages: tuple[PageAssessment, ...] = ()
    # OCR 需要栅格化的文件（如 Word 转出的 PDF）；None 表示直接对原始文件做 OCR。
    ocr_source: Path | None = None


def extract_text(path: Path) -> ExtractedText:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        with pymupdf.open(path) as document:
            pages = tuple(
                _assess_pdf_page(document[index], index)
                for index in range(document.page_count)
            )
        text = _normalize_text("\n".join(page.text for page in pages))
        return ExtractedText(
            text=text,
            page_count=len(pages),
            requires_ocr=any(page.needs_ocr for page in pages),
            pages=pages,
        )
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".doc":
        text = _normalize_text(_extract_legacy_doc(path))
        if not text:
            raise LegacyDocConversionError(
                "旧版 DOC 未提取到文字，请将文件另存为 DOCX 后重试"
            )
        return ExtractedText(text=text, page_count=1, requires_ocr=False)
    raise ValueError(f"Text extraction is unsupported for {suffix or 'unknown'}")


def _extract_docx(path: Path) -> ExtractedText:
    """提取 docx 正文（段落 + 表格 + 文本框）；正文过少时降级转 PDF 补齐。"""
    try:
        document = Document(path)
    except Exception:
        # 非标准 docx（如 .doc 改名或损坏）：降级 Word COM 提取。
        text = _normalize_text(_extract_legacy_doc(path))
        if not text:
            raise LegacyDocConversionError(
                "无法读取 Word 文档，请另存为 DOCX 后重试"
            )
        return ExtractedText(text=text, page_count=1, requires_ocr=False)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    parts.extend(_extract_textbox_paragraphs(document))
    text = _normalize_text("\n".join(parts))
    if analyze_text(text).valid_char_count >= MIN_MEANINGFUL_CHARS:
        return ExtractedText(text=text, page_count=1, requires_ocr=False)
    # 正文过少：关键信息可能只在文本框/图片/复杂排版中，转 PDF 复用文本层提取。
    try:
        pdf_path = convert_doc_to_pdf(path)
    except LegacyDocConversionError:
        return ExtractedText(text=text, page_count=1, requires_ocr=False)
    try:
        fallback = extract_text(pdf_path)
    except Exception:
        try:
            pdf_path.unlink(missing_ok=True)
        except OSError:
            pass
        return ExtractedText(text=text, page_count=1, requires_ocr=False)
    # 保留 PDF 供后续 OCR 栅格化；调用方用完负责清理 ocr_source。
    return ExtractedText(
        text=_normalize_text(fallback.text),
        page_count=fallback.page_count,
        requires_ocr=fallback.requires_ocr,
        pages=fallback.pages,
        ocr_source=pdf_path,
    )


def _extract_textbox_paragraphs(document) -> list[str]:
    """提取 Word 文本框（w:txbxContent）段落；python-docx 默认不读取。"""
    parts: list[str] = []
    for txbx in document.element.body.iter(qn("w:txbxContent")):
        for paragraph in txbx.iter(qn("w:p")):
            line = "".join(node.text or "" for node in paragraph.iter(qn("w:t")))
            if line.strip():
                parts.append(line)
    return parts


def _extract_legacy_doc(path: Path) -> str:
    system = platform.system()
    if system == "Windows":
        return _extract_doc_with_word(path)
    if system == "Darwin":
        return _extract_doc_with_textutil(path)
    raise LegacyDocConversionError(
        "当前系统不支持旧版 DOC，请将文件另存为 DOCX 后重试"
    )


def _extract_doc_with_word(path: Path) -> str:
    try:
        pythoncom = importlib.import_module("pythoncom")
        win32_client = importlib.import_module("win32com.client")
    except ModuleNotFoundError as error:
        raise LegacyDocConversionError(
            "Windows 解析旧版 DOC 需要 Microsoft Word，请安装 Word 或另存为 DOCX 后重试"
        ) from error

    word = None
    document = None
    com_initialized = False
    try:
        pythoncom.CoInitialize()
        com_initialized = True
        word = win32_client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(
            str(path.resolve()),
            ConfirmConversions=False,
            ReadOnly=True,
            AddToRecentFiles=False,
        )
        return str(document.Content.Text)
    except Exception as error:
        raise LegacyDocConversionError(
            "Microsoft Word 无法读取旧版 DOC，请将文件另存为 DOCX 后重试"
        ) from error
    finally:
        if document is not None:
            try:
                document.Close(SaveChanges=0)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        if com_initialized:
            pythoncom.CoUninitialize()


def convert_doc_to_pdf(source: Path) -> Path:
    """Convert a .doc/.docx to PDF via Microsoft Word for in-browser preview."""
    import tempfile
    import uuid

    try:
        pythoncom = importlib.import_module("pythoncom")
        win32_client = importlib.import_module("win32com.client")
    except ModuleNotFoundError as error:
        raise LegacyDocConversionError(
            "预览 Word 文档需要安装 Microsoft Word，请安装后重试"
        ) from error

    target = Path(tempfile.gettempdir()) / f"kerui_preview_{uuid.uuid4().hex}.pdf"
    word = None
    document = None
    com_initialized = False
    try:
        pythoncom.CoInitialize()
        com_initialized = True
        word = win32_client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(
            str(source.resolve()),
            ConfirmConversions=False,
            ReadOnly=True,
            AddToRecentFiles=False,
        )
        document.SaveAs2(str(target), FileFormat=17)  # wdFormatPDF
        document.Close(SaveChanges=0)
        document = None
        return target
    except Exception as error:
        raise LegacyDocConversionError(
            "Microsoft Word 无法将文档转换为 PDF 用于预览"
        ) from error
    finally:
        if document is not None:
            try:
                document.Close(SaveChanges=0)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        if com_initialized:
            pythoncom.CoUninitialize()


def _extract_doc_with_textutil(path: Path) -> str:
    try:
        result = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(path.resolve())],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=60,
            check=False,
        )
    except (FileNotFoundError, subprocess.TimeoutExpired) as error:
        raise LegacyDocConversionError(
            "macOS textutil 无法使用，请将旧版 DOC 另存为 DOCX 后重试"
        ) from error
    if result.returncode != 0:
        raise LegacyDocConversionError(
            "macOS textutil 无法读取旧版 DOC，请将文件另存为 DOCX 后重试"
        )
    return result.stdout


def _normalize_text(text: str) -> str:
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def _assess_pdf_page(page: pymupdf.Page, page_index: int) -> PageAssessment:
    """逐页判断直接提取的文本是否足以代表页面正文。

    仅把「空白/扫描、水印主导、重复乱码、图片为主且文字过少」的页面判为需要
    OCR，普通文本页（即使带头像、Logo、二维码或中文字符偏少）保持直接提取。
    """
    raw_text = page.get_text("text")
    text = _normalize_text(raw_text)
    quality = analyze_text(text)
    image_coverage = _image_coverage(page)

    needs_ocr = False
    reasons: list[str] = []
    if quality.valid_char_count == 0:
        needs_ocr = True
        reasons.append("页面无可提取文字")
    elif quality.dominant_ratio >= DOMINANT_FRAGMENT_RATIO:
        needs_ocr = True
        reasons.append(f"文字以重复水印为主（占 {quality.dominant_ratio:.0%}）")
    elif (
        quality.valid_char_count < MIN_MEANINGFUL_CHARS
        and image_coverage >= IMAGE_HEAVY_COVERAGE
    ):
        needs_ocr = True
        reasons.append(
            f"页面以图片为主且可提取文字过少（{quality.valid_char_count} 字符）"
        )

    return PageAssessment(
        page_index=page_index,
        text=text,
        needs_ocr=needs_ocr,
        reason="；".join(reasons) if needs_ocr else "文本正常，直接提取",
        valid_char_count=quality.valid_char_count,
        repeated_ratio=quality.repeated_ratio,
        dominant_ratio=quality.dominant_ratio,
        image_coverage=image_coverage,
    )


def _image_coverage(page: pymupdf.Page) -> float:
    """估算页面被图片覆盖的比例（按面积累加、上限 1.0）。"""
    rect = page.rect
    area = rect.width * rect.height
    if area <= 0:
        return 0.0
    infos = page.get_image_info()
    if not infos:
        return 0.0
    covered = 0.0
    for info in infos:
        x0, y0, x1, y1 = info["bbox"]
        width = max(0.0, min(x1, rect.x1) - max(x0, rect.x0))
        height = max(0.0, min(y1, rect.y1) - max(y0, rect.y0))
        covered += width * height
    return min(1.0, covered / area)
